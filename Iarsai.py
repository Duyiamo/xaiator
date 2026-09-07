import json
import sqlite3
from datetime import datetime
import io
import os
import folium
from folium.plugins import Draw, LocateControl
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st
from streamlit_folium import st_folium

DB_FILE = "quan_ly_dat_dai.db"


def init_db():
  conn = sqlite3.connect(DB_FILE)
  cursor = conn.cursor()
  cursor.execute("""
        CREATE TABLE IF NOT EXISTS thia_dat (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ho_ten TEXT,
            sdt TEXT,
            dia_chi_thuong_tru TEXT,
            thon_lang TEXT,
            so_to TEXT,
            so_thua TEXT,
            dia_chi_thua_dat TEXT,
            dien_tich_khai_bao REAL,
            nguon_goc TEXT,
            hien_trang TEXT,
            hien_trang_chi_tiet TEXT,
            tinh_trang_so TEXT,
            ghi_chu TEXT,
            lat REAL,
            lon REAL,
            geo_type TEXT,
            geo_coords TEXT,
            ngay_tao TEXT
        )
    """)
  try:
    cursor.execute("ALTER TABLE thia_dat ADD COLUMN thon_lang TEXT;")
  except sqlite3.OperationalError:
    pass
  try:
    cursor.execute("ALTER TABLE thia_dat ADD COLUMN geo_type TEXT;")
  except sqlite3.OperationalError:
    pass
  try:
    cursor.execute("ALTER TABLE thia_dat ADD COLUMN geo_coords TEXT;")
  except sqlite3.OperationalError:
    pass
  conn.commit()
  conn.close()


init_db()

st.set_page_config(
    page_title="Quản lý Hiện trạng Đất đai Cấp Xã", layout="wide"
)

st.title("🌾 Hệ thống Thu thập & Quản lý Hiện trạng Đất đai Cấp Xã")
st.markdown(
    "Ứng dụng hỗ trợ ghi nhận vị trí và ranh giới canh tác của hộ dân theo"
    " từng thôn, kết xuất báo cáo chuyên nghiệp."
)

tab1, tab2 = st.tabs(
    ["📝 1. Khai báo / Cập nhật thửa đất", "🔒 2. Khu vực Quản trị (Admin)"]
)

with tab1:
  st.header("Nhập thông tin và xác định vị trí / ranh giới thửa đất")

  with st.form("form_khai_bao"):
    col1, col2 = st.columns(2)
    with col1:
      ho_ten = st.text_input("Họ và tên chủ sử dụng *")
      sdt = st.text_input("Số điện thoại")
      dia_chi_thuong_tru = st.text_input(
          "Địa chỉ thường trú (Thôn/Xóm, Xã...)"
      )
      thon_lang = st.text_input(
          "Thôn / Làng tọa lạc thửa đất * (Ví dụ: Làng Hnáp, Thôn 1...)"
      )
      so_to = st.text_input("Số tờ bản đồ (nếu biết)")
      so_thua = st.text_input("Số thửa đất (nếu biết)")

    with col2:
      dia_chi_thua_dat = st.text_input(
          "Mô tả thêm khu vực thửa đất (Ví dụ: Khu Đồng Lớn, giáp suối...)"
      )
      dien_tich = st.number_input(
          "Diện tích tự khai báo (m²) *", min_value=0.0, value=0.0, step=10.0
      )
      nguon_goc = st.text_input(
          "Nguồn gốc sử dụng đất tự kê khai (Ví dụ: Khai hoang, Nhận chuyển"
          " nhượng...)"
      )
      hien_trang = st.selectbox(
          "Nhóm hiện trạng sử dụng đất *", [
              "Đất trồng lúa",
              "Đất trồng cây hàng năm khác",
              "Đất trồng cây lâu năm",
              "Đất nuôi trồng thủy sản",
              "Đất lâm nghiệp",
              "Đất phi nông nghiệp / Khác",
          ],
      )
      hien_trang_chi_tiet = st.text_input(
          "Cụ thể tên cây trồng / mục đích (Ví dụ: Cây mì, Cây điều, Cây keo,"
          " ...)"
      )
      tinh_trang_so = st.radio(
          "Tình trạng Giấy chứng nhận QSDĐ:",
          ["Đã có Giấy chứng nhận", "Chưa có / Đang sử dụng ổn định"],
          horizontal=True,
      )
      ghi_chu = st.text_area("Ghi chú thêm (nếu có)")

    st.markdown("---")
    st.markdown(
        "**Xác định vị trí trên bản đồ vệ tinh:** Bạn có thể **chấm 1 điểm"
        " (Marker)** vào giữa thửa đất hoặc dùng công cụ **vẽ đa giác/hình chữ"
        " nhật (Polygon/Rectangle)** để khoanh trọn ranh giới khu đất."
    )

    m = folium.Map(location=[13.304687, 108.603443], zoom_start=15)

    folium.TileLayer(
        tiles="https://{s}.google.com/vt/lyrs=s&x={x}&y={y}&z={z}",
        attr="Google Satellite",
        name="Bản đồ Vệ tinh",
        subdomains=["mt0", "mt1", "mt2", "mt3"],
        max_zoom=22,
        max_native_zoom=20,
        overlay=True,
        control=True,
    ).add_to(m)

    folium.LayerControl().add_to(m)

    LocateControl(
        auto_start=False,
        position="topleft",
        strings={
            "title": "Tìm vị trí hiện tại của tôi",
            "popup": "Bạn đang ở đây",
        },
        locate_options={"maxZoom": 18},
    ).add_to(m)

    draw = Draw(
        export=False,
        draw_options={
            "polyline": False,
            "polygon": True,
            "rectangle": True,
            "circle": False,
            "marker": True,
            "circlemarker": False,
        },
    )
    draw.add_to(m)

    output = st_folium(m, width="100%", height=450, key="map_input")

    submit_button = st.form_submit_button(
        "Gửi thông tin thửa đất", type="primary"
    )

    if submit_button:
      if not ho_ten or not thon_lang:
        st.error(
            "Vui lòng nhập đầy đủ [Họ và tên] và [Thôn / Làng tọa lạc thửa"
            " đất]!"
        )
      else:
        lat, lon = None, None
        geo_type = None
        geo_coords = None

        if output:
          if output.get("all_drawings") and len(output["all_drawings"]) > 0:
            last_shape = output["all_drawings"][-1]
            geometry = last_shape.get("geometry")
            if geometry:
              geo_type = geometry["type"]
              coords = geometry["coordinates"]
              geo_coords = json.dumps(coords)

              if geo_type == "Polygon" and len(coords) > 0 and len(coords[0]) > 0:
                pts = coords[0]
                lon = sum(pt[0] for pt in pts) / len(pts)
                lat = sum(pt[1] for pt in pts) / len(pts)
              elif geo_type == "Point" and len(coords) >= 2:
                lon = coords[0]
                lat = coords[1]

          if (
              (not lat or not lon)
              and output.get("last_clicked")
              and output["last_clicked"]
          ):
            lat = output["last_clicked"]["lat"]
            lon = output["last_clicked"]["lng"]
            geo_type = "Point"
            geo_coords = json.dumps([[lon, lat]])

        if not lat or not lon:
          st.error(
              "⚠️ Bạn chưa chấm điểm hoặc vẽ ranh giới thửa đất trên bản đồ! Vui"
              " lòng chọn vị trí trên bản đồ trước khi gửi."
          )
        else:
          conn = sqlite3.connect(DB_FILE)
          cursor = conn.cursor()

          is_duplicate = False
          cursor.execute(
              """
                    SELECT COUNT(*) FROM thia_dat 
                    WHERE ABS(lat - ?) < 0.00001 AND ABS(lon - ?) < 0.00001
                """,
              (lat, lon),
          )
          count = cursor.fetchone()[0]
          if count > 0:
            is_duplicate = True

          if is_duplicate:
            st.error(
                "⚠️ Vị trí thửa đất này đã được kê khai vào hệ thống! Xin vui"
                " lòng chọn vị trí khác trên bản đồ."
            )
          else:
            ngay_hien_tai = datetime.now().strftime("%Y-%m-%d")
            cursor.execute(
                """
                        INSERT INTO thia_dat (ho_ten, sdt, dia_chi_thuong_tru, thon_lang, so_to, so_thua, dia_chi_thua_dat, dien_tich_khai_bao, nguon_goc, hien_trang, hien_trang_chi_tiet, tinh_trang_so, ghi_chu, lat, lon, geo_type, geo_coords, ngay_tao)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                (
                    ho_ten,
                    sdt,
                    dia_chi_thuong_tru,
                    thon_lang,
                    so_to,
                    so_thua,
                    dia_chi_thua_dat,
                    dien_tich,
                    nguon_goc,
                    hien_trang,
                    hien_trang_chi_tiet,
                    tinh_trang_so,
                    ghi_chu,
                    lat,
                    lon,
                    geo_type,
                    geo_coords,
                    ngay_hien_tai,
                ),
            )
            conn.commit()
            st.success(
                f"Cảm ơn ông/bà **{ho_ten}**! Thông tin thửa đất tại **{thon_lang}**"
                " đã được gửi về hệ thống thành công."
            )

          conn.close()

with tab2:
  st.header("Khu vực Quản trị dành cho Cán bộ địa chính xã Ia RSai")

  password = st.text_input(
      "Nhập mật khẩu quản lý để tiếp tục:", type="password"
  )
  ADMIN_PASSWORD = "phuc123"

  if password == ADMIN_PASSWORD:
    st.success("Xác thực thành công! Chào mừng cán bộ quản lý.")

    conn = sqlite3.connect(DB_FILE)
    df = pd.read_sql_query("SELECT * FROM thia_dat", conn)
    conn.close()

    if df.empty:
      st.info("Chưa có dữ liệu khai báo nào từ người dân.")
    else:
      st.metric(
          label="Tổng số thửa đã cập nhật vào hệ thống",
          value=f"{len(df)} thửa đất",
      )

      st.markdown("### Bộ lọc dữ liệu quản lý theo không gian hành chính")

      col_f1, col_f2 = st.columns(2)
      with col_f1:
        danh_sach_thon = ["Tất cả các Thôn/Làng"] + sorted(
            df["thon_lang"].dropna().unique().tolist()
        )
        chon_thon = st.selectbox("Lọc theo Thôn / Làng:", danh_sach_thon)

      with col_f2:
        danh_sach_ngay = ["Tất cả các ngày"] + sorted(
            df["ngay_tao"].dropna().unique().tolist()
        )
        chon_ngay = st.selectbox("Lọc theo ngày kê khai:", danh_sach_ngay)

      df_hien_thi = df.copy()
      if chon_thon != "Tất cả các Thôn/Làng":
        df_hien_thi = df_hien_thi[df_hien_thi["thon_lang"] == chon_thon]
      if chon_ngay != "Tất cả các ngày":
        df_hien_thi = df_hien_thi[df_hien_thi["ngay_tao"] == chon_ngay]

      st.markdown(f"Đang hiển thị **{len(df_hien_thi)}** bản ghi.")

      st.dataframe(
          df_hien_thi[[
              "id",
              "ho_ten",
              "sdt",
              "thon_lang",
              "so_to",
              "so_thua",
              "dia_chi_thua_dat",
              "dien_tich_khai_bao",
              "nguon_goc",
              "hien_trang",
              "tinh_trang_so",
              "ngay_tao",
          ]],
          use_container_width=True,
      )

      st.markdown("### 🔍 Kiểm tra nhanh vị trí / ranh giới từng thửa đất")
      if not df_hien_thi.empty:
        options_thua = []
        for _, r in df_hien_thi.iterrows():
          label_item = (
              f"ID: {r['id']} | Thôn: {r['thon_lang']} | Chủ hộ:"
              f" {r['ho_ten']} | Thửa: {r['so_thua']} - Tờ: {r['so_to']}"
          )
          options_thua.append((label_item, r["id"]))

        chon_lua = st.selectbox(
            "Chọn thửa đất cần kiểm tra trên bản đồ:",
            options_thua,
            format_func=lambda x: x[0],
        )

        if chon_lua:
          selected_id = chon_lua[1]
          row_chon = df_hien_thi[df_hien_thi["id"] == selected_id].iloc[0]

          lat_Check = row_chon["lat"]
          lon_Check = row_chon["lon"]

          if pd.notnull(lat_Check) and pd.notnull(lon_Check):
            m_admin = folium.Map(
                location=[lat_Check, lon_Check], zoom_start=18
            )

            folium.TileLayer(
                tiles="https://{s}.google.com/vt/lyrs=s&x={x}&y={y}&z={z}",
                attr="Google Satellite",
                name="Bản đồ Vệ tinh",
                subdomains=["mt0", "mt1", "mt2", "mt3"],
                max_zoom=22,
                max_native_zoom=20,
                overlay=True,
                control=True,
            ).add_to(m_admin)

            if (
                row_chon["geo_type"] == "Polygon"
                and pd.notnull(row_chon["geo_coords"])
            ):
              try:
                coords = json.loads(row_chon["geo_coords"])
                if len(coords) > 0:
                  folium_pts = [[pt[1], pt[0]] for pt in coords[0]]
                  folium.Polygon(
                      locations=folium_pts,
                      color="yellow",
                      weight=3,
                      fill=True,
                      fill_color="blue",
                      fill_opacity=0.3,
                      popup=f"<b>{row_chon['ho_ten']}</b><br>Thôn: {row_chon['thon_lang']}<br>Diện tích: {row_chon['dien_tich_khai_bao']} m²",
                  ).add_to(m_admin)
              except Exception:
                pass

            folium.Marker(
                [lat_Check, lon_Check],
                popup=(
                    f"<b>Chủ hộ: {row_chon['ho_ten']}</b><br>Thôn:"
                    f" {row_chon['thon_lang']}<br>Diện tích:"
                    f" {row_chon['dien_tich_khai_bao']} m²<br>Hiện trạng:"
                    f" {row_chon['hien_trang']} ({row_chon['hien_trang_chi_tiet']})"
                ),
                icon=folium.Icon(color="red", icon="home"),
            ).add_to(m_admin)

            st_folium(m_admin, width="100%", height=400, key=f"map_{selected_id}")
          else:
            st.warning("Thửa đất này chưa có thông tin vị trí trên bản đồ.")

      st.markdown("### Xuất dữ liệu phục vụ nội nghiệp")

      col_ex1, col_ex2 = st.columns(2)

      with col_ex1:
        if st.button("📥 Tạo và Tải xuống File Excel Báo Cáo"):
          wb = openpyxl.Workbook()
          ws = wb.active
          ws.title = "Danh sách hiện trạng đất"
          ws.sheet_view.showGridLines = True

          ws.merge_cells("A1:Q1")
          ws["A1"] = (
              "DANH SÁCH TỔNG HỢP HIỆN TRẠNG CANH TÁC ĐẤT ĐAI CẤP Xã"
          ).upper()
          ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
          ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

          headers = [
              "STT",
              "Họ và tên chủ sử dụng",
              "Số điện thoại",
              "Địa chỉ thường trú",
              "Thôn / Làng",
              "Số tờ",
              "Số thửa",
              "Địa chỉ thửa đất",
              "Diện tích (m²)",
              "Nguồn gốc tự kê khai",
              "Nhóm hiện trạng",
              "Tên cây trồng cụ thể",
              "Tình trạng Giấy chứng nhận",
              "Kiểu dữ liệu bản đồ",
              "Link Google Maps",
              "Ngày kê khai",
          ]
          ws.append([])
          ws.append(headers)

          header_font = Font(
              name="Times New Roman", size=11, bold=True, color="FFFFFF"
          )
          header_fill = PatternFill(
              start_color="1F4E78", end_color="1F4E78", fill_type="solid"
          )
          header_align = Alignment(
              horizontal="center", vertical="center", wrap_text=True
          )

          for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

          thin_border = Border(
              left=Side(style="thin", color="D9D9D9"),
              right=Side(style="thin", color="D9D9D9"),
              top=Side(style="thin", color="D9D9D9"),
              bottom=Side(style="thin", color="D9D9D9"),
          )
          data_font = Font(name="Times New Roman", size=11)

          for idx, row in df_hien_thi.reset_index(drop=True).iterrows():
            lat_val = row["lat"]
            lon_val = row["lon"]
            map_link = (
                f"https://www.google.com/maps?q={lat_val},{lon_val}"
                if pd.notnull(lat_val) and pd.notnull(lon_val)
                else "Chưa có vị trí"
            )

            row_data = [
                idx + 1,
                row["ho_ten"],
                str(row["sdt"]) if pd.notnull(row["sdt"]) else "",
                str(row["dia_chi_thuong_tru"])
                if pd.notnull(row["dia_chi_thuong_tru"])
                else "",
                str(row["thon_lang"]) if pd.notnull(row["thon_lang"]) else "",
                str(row["so_to"]) if pd.notnull(row["so_to"]) else "",
                str(row["so_thua"]) if pd.notnull(row["so_thua"]) else "",
                str(row["dia_chi_thua_dat"])
                if pd.notnull(row["dia_chi_thua_dat"])
                else "",
                row["dien_tich_khai_bao"],
                str(row["nguon_goc"]) if pd.notnull(row["nguon_goc"]) else "",
                row["hien_trang"],
                row["hien_trang_chi_tiet"]
                if pd.notnull(row["hien_trang_chi_tiet"])
                else "",
                row["tinh_trang_so"],
                row["geo_type"] if pd.notnull(row["geo_type"]) else "",
                map_link,
                str(row["ngay_tao"]),
            ]
            ws.append(row_data)

          for row_idx in range(4, 4 + len(df_hien_thi)):
            for col_idx in range(1, len(headers) + 1):
              cell = ws.cell(row=row_idx, column=col_idx)
              cell.font = data_font
              cell.border = thin_border

              if col_idx in [1, 6, 7, 14, 16]:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
              elif col_idx == 9:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
              elif col_idx == 15:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                cell.font = Font(
                    name="Times New Roman",
                    size=10,
                    color="0563C1",
                    underline="single",
                )
              else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

          for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
              if cell.row > 1:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                  max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

          local_file_name = "Bao_cao_hien_trang_dat_dai_theo_thon.xlsx"
          wb.save(local_file_name)

          with open(local_file_name, "rb") as f:
            st.download_button(
                label="📥 Tải file Excel ngay",
                data=f,
                file_name=local_file_name,
                mime=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            )

      with col_ex2:
        if st.button("🌍 Tải File Bản Đồ (KML hỗ trợ Vùng & Điểm)"):
          kml_content = """<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
  <Document>
    <name>Danh sách Thửa đất theo Thôn</name>
"""
          for _, row in df_hien_thi.iterrows():
            ho_ten_Val = str(row["ho_ten"] or "").strip()
            thon_val = str(row["thon_lang"] or "").strip()
            so_to_val = str(row["so_to"] or "").strip()
            so_thua_val = str(row["so_thua"] or "").strip()

            if so_to_val != "" and so_thua_val != "":
              name = (
                  f"[{thon_val}] {ho_ten_Val} - Thửa: {so_thua_val}, Tờ:"
                  f" {so_to_val}"
              )
            else:
              name = f"[{thon_val}] {ho_ten_Val}"

            desc = (
                f"Thôn: {thon_val}<br/>Chủ sử dụng:"
                f" {ho_ten_Val}<br/>SĐT: {row['sdt']}<br/>Diện tích khai"
                f" báo: {row['dien_tich_khai_bao']} m²<br/>Hiện trạng:"
                f" {row['hien_trang']} ({row['hien_trang_chi_tiet']})"
            )

            geo_type = row["geo_type"]
            geo_coords_str = row["geo_coords"]

            if geo_type == "Polygon" and pd.notnull(geo_coords_str):
              try:
                coords_list = json.loads(geo_coords_str)
                if len(coords_list) > 0:
                  kml_coords_flat = " ".join(
                      [f"{pt[0]},{pt[1]},0" for pt in coords_list[0]]
                  )
                  kml_content += f"""    <Placemark>
      <name><![CDATA[{name}]]></name>
      <description><![CDATA[{desc}]]></description>
      <Polygon>
        <outerBoundaryIs>
          <LinearRing>
            <coordinates>{kml_coords_flat}</coordinates>
          </LinearRing>
        </outerBoundaryIs>
      </Polygon>
    </Placemark>
"""
              except Exception:
                pass

            elif pd.notnull(row["lat"]) and pd.notnull(row["lon"]):
              kml_content += f"""    <Placemark>
      <name><![CDATA[{name}]]></name>
      <description><![CDATA[{desc}]]></description>
      <Point>
        <coordinates>{row['lon']},{row['lat']},0</coordinates>
      </Point>
    </Placemark>
"""
          kml_content += """  </Document>
</kml>"""

          st.download_button(
              label="📥 Tải file KML mở Google Earth",
              data=kml_content.encode("utf-8"),
              file_name="hien_trang_dat_dai_theo_thon.kml",
              mime="application/vnd.google-earth.kml+xml",
          )

  elif password != "":
    st.error("Sai mật khẩu quản lý! Vui lòng thử lại.")
  else:
    st.info("Vui lòng nhập mật khẩu quản lý để xem danh sách và xuất báo cáo.")
